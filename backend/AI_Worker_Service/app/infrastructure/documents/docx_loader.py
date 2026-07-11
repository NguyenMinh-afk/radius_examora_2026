"""
DOCX document loader using python-docx.
"""
import io

from app.core.exceptions import DocumentError
from app.core.logging import get_logger

logger = get_logger(__name__)


def load_docx(file_bytes: bytes, filename: str = "file.docx") -> str:
    """
    Extract text from a DOCX file, preserving paragraph structure.

    Args:
        file_bytes: Raw DOCX bytes
        filename: Original filename for logging

    Returns:
        Extracted text content

    Raises:
        DocumentError: If DOCX cannot be read or is empty
    """
    try:
        from docx import Document
    except ImportError:
        raise DocumentError("python-docx is not installed. Run: pip install python-docx")

    try:
        doc = Document(io.BytesIO(file_bytes))

        paragraphs = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                paragraphs.append(text)

        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    paragraphs.append(" | ".join(row_texts))

        full_text = "\n".join(paragraphs).strip()

        if not full_text:
            raise DocumentError(f"DOCX '{filename}' is empty or contains no readable text.")

        logger.info(
            "Loaded DOCX '%s': %d paragraphs, %d characters",
            filename,
            len(paragraphs),
            len(full_text),
        )
        return full_text

    except DocumentError:
        raise
    except Exception as e:
        raise DocumentError(f"Failed to read DOCX '{filename}': {e}") from e
